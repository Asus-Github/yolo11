#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
import re
import shlex
import subprocess
import sys
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

LOCAL_ROOT = Path("/Users/asus/ultralytics")
EXCEL_PATH = LOCAL_ROOT / "刘华硕-飞书导入实验记录表.xlsx"
REMOTE = "autodl"
REMOTE_ROOT = "/root/autodl-tmp/ultralytics"


def run(cmd: list[str], *, check: bool = True, capture: bool = False, timeout: int | None = None) -> subprocess.CompletedProcess[str]:
    print("[local] $", " ".join(shlex.quote(c) for c in cmd), flush=True)
    return subprocess.run(cmd, text=True, check=check, capture_output=capture, timeout=timeout)


def ssh(command: str, *, check: bool = True, capture: bool = False, timeout: int | None = None) -> subprocess.CompletedProcess[str]:
    return run(["ssh", REMOTE, command], check=check, capture=capture, timeout=timeout)


def scp_from(remote: str, local: Path) -> None:
    local.parent.mkdir(parents=True, exist_ok=True)
    if local.exists() and local.is_dir():
        backup = local.with_name(local.name + f".backup-{int(time.time())}")
        local.rename(backup)
        print(f"[sync] moved existing {local} to {backup}", flush=True)
    run(["scp", "-r", f"{REMOTE}:{remote}", str(local)], timeout=600)


def remote_alive() -> bool:
    try:
        ssh("echo ok", timeout=30, capture=True)
        return True
    except Exception as e:
        print(f"[warn] ssh unavailable: {e}", flush=True)
        return False


def tmux_exists(session: str) -> bool:
    cp = ssh(f"tmux has-session -t {shlex.quote(session)} 2>/dev/null", check=False, timeout=30)
    return cp.returncode == 0


def d3_done() -> bool:
    cmd = f"cd {REMOTE_ROOT} && test -f runs/detect/runs/ablation/TDW-CO-D3/results.csv && test -f runs/detect/runs/ablation/TDW-CO-D3/weights/best.pt && grep -Eq 'epochs completed|EarlyStopping' runs/TDW-CO-D3_train.log"
    cp = ssh(cmd, check=False, timeout=60)
    return cp.returncode == 0


def print_d3_progress() -> None:
    cmd = f"cd {REMOTE_ROOT} && /root/miniconda3/bin/python - <<'PY'\nfrom pathlib import Path\np=Path('runs/detect/runs/ablation/TDW-CO-D3/results.csv')\nif p.exists():\n    lines=p.read_text().strip().splitlines()\n    print('results_rows', max(0, len(lines)-1))\nlog=Path('runs/TDW-CO-D3_train.log')\nif log.exists():\n    txt=log.read_text(errors='ignore')\n    import re\n    epochs=re.findall(r'\\n\\s*(\\d+)/300\\s+', txt)\n    if epochs: print('last_epoch_seen', epochs[-1])\n    for line in txt.splitlines()[-20:]:\n        if 'all       1411' in line or 'epochs completed' in line or 'EarlyStopping' in line:\n            print(line)\nPY"
    ssh(cmd, check=False, timeout=60)


def parse_remote(exp: str) -> dict[str, Any]:
    cmd = f"cd {REMOTE_ROOT} && /root/miniconda3/bin/python - <<'PY'\nimport csv, json, re\nfrom pathlib import Path\nexp={exp!r}\np=Path('runs/detect/runs/ablation')/exp/'results.csv'\nrows=list(csv.DictReader(p.open()))\ndef get(r,n):\n    for k,v in r.items():\n        if k.strip()==n: return v\nbest=max(rows, key=lambda r: float(get(r,'metrics/mAP50-95(B)') or 0))\nlog=Path('runs')/f'{exp}_train.log'\ntxt=log.read_text(errors='replace')\nlast=''\nfor line in txt.splitlines():\n    if 'summary (fused)' in line or 'summary:' in line:\n        last=line\nparams_m=None; gflops=None\nm=re.search(r'([0-9,]+) parameters.*?([0-9.]+) GFLOPs', last)\nif m:\n    params_m=round(int(m.group(1).replace(',', ''))/1_000_000, 2)\n    gflops=round(float(m.group(2)), 2)\nout=dict(name=exp, epochs_recorded=len(rows), best_epoch=int(float(get(best,'epoch'))), precision=float(get(best,'metrics/precision(B)')), recall=float(get(best,'metrics/recall(B)')), map50=float(get(best,'metrics/mAP50(B)')), map50_95=float(get(best,'metrics/mAP50-95(B)')), params_m=params_m, gflops=gflops, wiou_logged='[ducc_wiou_train] BboxLoss.iou_type = wiou' in txt)\nprint(json.dumps(out, ensure_ascii=False))\nPY"
    cp = ssh(cmd, capture=True, timeout=120)
    return json.loads((cp.stdout or "").strip().splitlines()[-1])


def update_excel_d3(summary: dict[str, Any]) -> None:
    wb = load_workbook(EXCEL_PATH)
    ws = wb["数据表"]
    row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value == "TDW-CO-D3" and ws.cell(r, 5).value == "WIoU":
            row = r
            break
    if row is None:
        ws.append([None, "下一阶段消融实验", "TDW-CO-D3", "yolo11-tdw-co-d3.yaml", "WIoU", "✅", "✅", "✅"])
        row = ws.max_row
    ws.cell(row, 9).value = round(summary["precision"] * 100, 1)
    ws.cell(row, 10).value = round(summary["recall"] * 100, 1)
    ws.cell(row, 11).value = round(summary["map50"] * 100, 1)
    ws.cell(row, 12).value = round(summary["map50_95"] * 100, 1)
    if summary.get("gflops") is not None:
        ws.cell(row, 14).value = summary["gflops"]
    if summary.get("params_m") is not None:
        ws.cell(row, 15).value = summary["params_m"]
    wiou = "WIoU已确认" if summary.get("wiou_logged") else "WIoU日志未确认"
    ws.cell(row, 16).value = f"best epoch={summary['best_epoch']}；{summary['epochs_recorded']} epochs recorded；{wiou}；AutoDL训练完成并同步本地。"
    wb.save(EXCEL_PATH)
    print(f"[excel] updated TDW-CO-D3 row {row}: {summary}", flush=True)


def sync_d3() -> None:
    scp_from(f"{REMOTE}:{REMOTE_ROOT}/runs/detect/runs/ablation/TDW-CO-D3".split(":", 1)[1], LOCAL_ROOT / "runs/detect/runs/ablation/TDW-CO-D3")
    scp_from(f"{REMOTE_ROOT}/runs/TDW-CO-D3_train.log", LOCAL_ROOT / "runs/TDW-CO-D3_train.log")


def commit_push_d3_nonfatal() -> None:
    cmd = f"""
set -euo pipefail
source /etc/network_turbo
cd {REMOTE_ROOT}
git checkout feat/tdw-co-d3
git add -f runs/TDW-CO-D3_train.log runs/detect/runs/ablation/TDW-CO-D3
git reset -q -- runs/detect/runs/ablation/TDW-CO-D3/weights || true
if git diff --cached --quiet; then
  echo '[git] no non-weight D3 artifacts to commit'
else
  git commit -m 'feat(results): add TDW-CO-D3 WIoU training results' || true
  git push origin feat/tdw-co-d3 || true
fi
"""
    ssh(cmd, check=False, timeout=600)


def start_dwca() -> None:
    cmd = f"""
set -euo pipefail
source /etc/network_turbo
cd {REMOTE_ROOT}
git fetch origin
git checkout feat/dwca-p2
git pull --ff-only origin feat/dwca-p2
(tmux has-session -t t_train_dwca_p2 2>/dev/null && tmux kill-session -t t_train_dwca_p2) || true
tmux new-session -d -s t_train_dwca_p2
tmux set-option -t t_train_dwca_p2 remain-on-exit on
tmux set-option -g mouse on
tmux set-window-option -g history-limit 100000
tmux send-keys -t t_train_dwca_p2 "source /etc/network_turbo && cd {REMOTE_ROOT} && /root/miniconda3/bin/python /root/autodl-tmp/ducc_wiou_train.py --cfg ultralytics/cfg/models/11/yolo11-dwca-p2.yaml --name DWCA-P2 --epochs 300 --batch 0.85 --imgsz 640 --seed 42 --patience 50 2>&1 | tee runs/DWCA-P2_train.log" Enter
echo started_t_train_dwca_p2
"""
    ssh(cmd, timeout=180)


def main() -> int:
    print("[monitor] watching TDW-CO-D3; will start DWCA-P2 after completion", flush=True)
    while True:
        if not remote_alive():
            time.sleep(120)
            continue
        print_d3_progress()
        if d3_done():
            print("[monitor] TDW-CO-D3 appears complete", flush=True)
            summary = parse_remote("TDW-CO-D3")
            print(f"[monitor] D3 summary: {summary}", flush=True)
            sync_d3()
            update_excel_d3(summary)
            commit_push_d3_nonfatal()
            start_dwca()
            print("[monitor] DWCA-P2 started", flush=True)
            return 0
        time.sleep(300)


if __name__ == "__main__":
    raise SystemExit(main())
