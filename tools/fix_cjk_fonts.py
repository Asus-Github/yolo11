from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path("docs/paper_outputs/中文版北核_DWCA-P2-哈理工学报格式版_格式修订.docx")
OUTPUT = Path("docs/paper_outputs/中文版北核_DWCA-P2-哈理工学报格式版_最终格式修订.docx")
CJK = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")


def set_east_asia(run, font_name):
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


document = Document(SOURCE)
for para in document.paragraphs:
    is_heading = para.style.name in {"Heading 1", "Heading 2"}
    is_caption = para.text.strip().startswith(("图", "表"))
    for run in para.runs:
        if CJK.search(run.text):
            set_east_asia(run, "黑体" if is_heading or is_caption else "宋体")
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if CJK.search(run.text):
                        set_east_asia(run, "宋体")
document.save(OUTPUT)
print(OUTPUT.resolve())
