from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path("docs/paper_outputs/中文版北核_DWCA-P2-哈理工学报格式版_图8空白修复.docx")
OUTPUT = Path("docs/paper_outputs/中文版北核_DWCA-P2-哈理工学报格式版_格式修订.docx")


def replace_literal_page_number(paragraph):
    if paragraph.text.strip() != "1":
        return
    for child in list(paragraph._p):
        paragraph._p.remove(child)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    displayed = OxmlElement("w:t")
    displayed.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, displayed, end])


document = Document(SOURCE)
for section in document.sections:
    for paragraph in section.footer.paragraphs:
        replace_literal_page_number(paragraph)
document.save(OUTPUT)
print(OUTPUT.resolve())
